import openpyxl
import os
from openpyxl.styles import Alignment 
from openpyxl.drawing.image import Image as OpenpyxlImage
from datetime import datetime, date

# --- 1. FUNÇÃO PARA INSERIR A LOGO FIXA ---
def inserir_logo_fixa(sheet, celula_ancora='A3'):
    """
    Busca 'logo_oficial.png' na mesma pasta do script e insere na planilha.
    """
    try:
        # Pega o diretório onde este arquivo .py está
        base_dir = os.path.dirname(os.path.abspath(__file__))
        imagem_path = os.path.join(base_dir, "logo_oficial.png")

        if os.path.exists(imagem_path):
            img = OpenpyxlImage(imagem_path)
            # Ajuste o tamanho se necessário
            # img.width = 120
            # img.height = 120
            sheet.add_image(img, celula_ancora)
            return True
        else:
            print(f"⚠️ Aviso: A imagem {imagem_path} não foi encontrada. A ficha ficará sem logo.")
            return False
    except Exception as e:
        print(f"❌ Erro ao inserir logo: {e}")
        return False

# --- 2. FORMATAÇÃO DE DATA (BR) ---
def formatar_valor_para_br(valor):
    if valor is None:
        return ""
    if isinstance(valor, (date, datetime)):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, str):
        valor = valor.strip()
        # Detecta formato YYYY-MM-DD
        if len(valor) == 10 and valor[4] == '-' and valor[7] == '-':
            try:
                data_obj = datetime.strptime(valor, '%Y-%m-%d')
                return data_obj.strftime('%d/%m/%Y')
            except ValueError:
                pass
    return str(valor)

def formatar_data_por_extenso(data_obj):
    if not isinstance(data_obj, (datetime, date)):
        try:
            if isinstance(data_obj, str) and len(data_obj) == 10:
                 data_obj = datetime.strptime(data_obj, '%Y-%m-%d')
            else:
                return "data não informada"
        except:
            return "data não informada"
    meses = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
    return f"{data_obj.day} de {meses[data_obj.month - 1]} de {data_obj.year}"

# --- 3. SUBSTITUIÇÃO INTELIGENTE (Preenche e Limpa) ---
def substituir_placeholders_dinamicos(sheet, lista_dados, mapeamento_chaves, max_itens):
    valores_para_escrever = {}
    
    for i in range(max_itens):
        numero_excel = i + 1
        
        if i < len(lista_dados):
            item = lista_dados[i]
            for prefixo_excel, chave_banco in mapeamento_chaves.items():
                placeholder = f"{prefixo_excel}{numero_excel}"
                valor_bruto = item.get(chave_banco, "")
                valores_para_escrever[placeholder] = formatar_valor_para_br(valor_bruto)
        else:
            # Limpa as linhas excedentes
            for prefixo_excel in mapeamento_chaves.keys():
                placeholder = f"{prefixo_excel}{numero_excel}"
                valores_para_escrever[placeholder] = ""

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                valor_celula = cell.value.strip()
                if valor_celula in valores_para_escrever:
                    cell.value = valores_para_escrever[valor_celula]
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# --- 4. FUNÇÃO PRINCIPAL ---
def preencher_ficha_excel(template_path, dados_funcionario, caminho_saida):
    try:
        workbook = openpyxl.load_workbook(template_path)
        estilo_centralizado = Alignment(horizontal='center', vertical='center', wrap_text=True)
        estilo_esquerda = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # ================= ABA FRENTE =================
        sheet_frente = workbook["FRENTE"]
        sheet_frente.page_setup.orientation = 'landscape'
        sheet_frente.page_setup.paperSize = 9 # A4
        sheet_frente.page_setup.fitToPage = True

        # >>> AQUI INSERIMOS A LOGO MANUALMENTE <<<
        inserir_logo_fixa(sheet_frente, celula_ancora='A3')

        mapeamento_frente = {
            'AA7': dados_funcionario.get('matricula'), 'F13': dados_funcionario.get('nome'),
            'B19': dados_funcionario.get('data_nascimento'), 'F19': dados_funcionario.get('estado_civil'),
            'I19': dados_funcionario.get('sexo'), 'M19': dados_funcionario.get('naturalidade'),
            'R19': dados_funcionario.get('nacionalidade'), 'V19': dados_funcionario.get('carteira_profissional'),
            'B22': dados_funcionario.get('servico_militar'), 'F22': dados_funcionario.get('titulo_eleitor'),
            'I22': dados_funcionario.get('cpf'), 'M22': dados_funcionario.get('identidade'),
            'R22': dados_funcionario.get('pis'), 'V22': dados_funcionario.get('carteira_saude'),
            'F16': dados_funcionario.get('campo_mudança_nome'), 'D26': dados_funcionario.get('nome_pai'),
            'D27': dados_funcionario.get('nome_mae'), 'D30': dados_funcionario.get('endereco'),
            'B43': dados_funcionario.get('data_Admissao'), 'F43': dados_funcionario.get('data_posse'),
            'I43': dados_funcionario.get('cargo'), 'M43': dados_funcionario.get('vencimento_ou_salario'),
            'I46': dados_funcionario.get('horario'), 'B46': dados_funcionario.get('data_desligamento'),
            'F46': dados_funcionario.get('inicio_atividades'), 'M46': dados_funcionario.get('descanso_semanal'),
        }
        
        celulas_de_data_frente = ['B19', 'B43', 'F43', 'F46', 'B46']

        for celula_str, valor in mapeamento_frente.items():
            if celula_str in celulas_de_data_frente:
                val_str = formatar_valor_para_br(valor)
                if not val_str: val_str = "não informado"
            else:
                 val_str = str(valor) if (valor is not None and str(valor).strip() != '') else "não informado"

            sheet_frente[celula_str] = val_str
            align = estilo_esquerda if celula_str in ['F13', 'F16'] else estilo_centralizado
            sheet_frente[celula_str].alignment = align

        # Beneficiários
        beneficiarios = dados_funcionario.get('beneficiarios', [])
        tabelas_beneficiarios = [
            {'colunas': {'nome': 'C', 'parentesco': 'K', 'nascimento': 'L'}, 'start_index': 0},
            {'colunas': {'nome': 'P', 'parentesco': 'AA', 'nascimento': 'AB'}, 'start_index': 7}
        ]
        
        for tabela in tabelas_beneficiarios:
            linha_inicial = 34
            max_linhas = 7
            for i in range(max_linhas):
                indice = tabela['start_index'] + i
                if indice < len(beneficiarios):
                    ben = beneficiarios[indice]
                    dt_fmt = formatar_valor_para_br(ben.get('data_nascimento'))
                    
                    sheet_frente[f"{tabela['colunas']['nome']}{linha_inicial + i}"] = ben.get('nome', 'não informado')
                    sheet_frente[f"{tabela['colunas']['parentesco']}{linha_inicial + i}"] = ben.get('parentesco', 'não informado')
                    sheet_frente[f"{tabela['colunas']['nascimento']}{linha_inicial + i}"] = dt_fmt
                else:
                    sheet_frente[f"{tabela['colunas']['nome']}{linha_inicial + i}"] = ""
                    sheet_frente[f"{tabela['colunas']['parentesco']}{linha_inicial + i}"] = ""
                    sheet_frente[f"{tabela['colunas']['nascimento']}{linha_inicial + i}"] = ""

        # ================= ABA VERSO =================
        sheet_verso = workbook["VERSO"]
        sheet_verso.page_setup.orientation = 'landscape'
        sheet_verso.page_setup.paperSize = 9
        sheet_verso.page_setup.fitToPage = True

        # Transferências
        transferencias = dados_funcionario.get("transferencias", [])
        mapa_transf = {
            'CAMPO_DATA_TRASFERENCIA_': 'data_transferencia',
            'CAMPO_DE_LOTACAO_': 'de_lotacao',
            'CAMPO_PARA_LOTACAO_': 'para_lotacao'
        }
        substituir_placeholders_dinamicos(sheet_verso, transferencias, mapa_transf, 19)

        # Férias
        ferias = dados_funcionario.get("ferias", [])
        mapa_ferias = {
            'CAMPO_PERIODO_AQUISICAO_': 'periodo_aquisicao',
            'CAMPO_PERIODO_UTILIZACAO_': 'periodo_utilizacao'
        }
        substituir_placeholders_dinamicos(sheet_verso, ferias, mapa_ferias, 11)

        # ================= ABA ANOTAÇÕES =================
        sheet_anotacoes = workbook["OUTRAS ANOTAÇÕES"]
        sheet_anotacoes.page_setup.orientation = 'landscape'
        sheet_anotacoes.page_setup.paperSize = 9
        sheet_anotacoes.page_setup.fitToPage = True

        dt_adm = formatar_data_por_extenso(dados_funcionario.get('data_Admissao'))
        dt_pub = formatar_data_por_extenso(dados_funcionario.get('data_publicacao'))

        map_anotacoes = {
            'A2': dados_funcionario.get('nome'),
            'B4': dados_funcionario.get('cargo'),
            'C3': dt_pub.upper(),
            'G3': dt_adm.upper(),
        }

        for celula, valor in map_anotacoes.items():
            val = str(valor) if valor is not None else "não informado"
            sheet_anotacoes[celula] = val
            sheet_anotacoes[celula].alignment = estilo_esquerda

        anotacoes = dados_funcionario.get("outras_anotacoes", "")
        sheet_anotacoes["A6"] = anotacoes
        sheet_anotacoes["A6"].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        workbook.save(caminho_saida)
        return True, None

    except KeyError as e:
        return False, f"Aba '{e}' não encontrada no template."
    except Exception as e:
        import traceback
        traceback.print_exc()
        return False, str(e)